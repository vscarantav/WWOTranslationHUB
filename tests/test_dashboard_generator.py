import os
import shutil
import sys
import unittest
import uuid
from pathlib import Path

from openpyxl import load_workbook


APP_DIR = Path(__file__).resolve().parents[1] / "app"
sys.path.insert(0, str(APP_DIR))

from core.dashboard_generator import DashboardGenerator


class DashboardGeneratorReviewSheetTests(unittest.TestCase):
    def test_report_identity_can_be_overridden_for_edtech(self):
        hub_dir = Path.cwd() / "tests" / f"_dashboard_test_{uuid.uuid4().hex}"
        hub_dir.mkdir(parents=True)
        try:
            lesson_path = hub_dir / "raw_html" / "lesson_1.html"
            lesson_path.parent.mkdir(parents=True)
            lesson_path.write_text("<p>Lesson content</p>", encoding="utf-8")
            log_path = hub_dir / "translation_log.txt"
            log_path.write_text(
                "\n".join([
                    "--- New Session (Target: PTBR) ---",
                    f"[2026-09-18 10:00:00] [System] TranslatedPage: First Lesson | {lesson_path}",
                ]),
                encoding="utf-8",
            )

            report_path = DashboardGenerator(
                str(log_path), str(hub_dir), "PTBR"
            ).generate(
                lambda _message: None,
                report_name="Writing Basics EdTech Master",
                report_code="EDTECH-PTBR",
                review_pages=[
                    {"title": "First Lesson", "filepath": str(lesson_path)},
                    {"title": "Second Lesson", "filepath": str(hub_dir / "raw_html" / "lesson_2.html")},
                ],
                excluded_sheets=("Dashboard", "Bot Analysis", "Raw Logs"),
            )

            self.assertEqual(
                Path(report_path).name,
                "Writing Basics EdTech Master Translation Report.xlsx",
            )
            workbook = load_workbook(report_path, data_only=False)
            self.assertNotIn("Dashboard", workbook.sheetnames)
            self.assertNotIn("Bot Analysis", workbook.sheetnames)
            self.assertNotIn("Raw Logs", workbook.sheetnames)
            self.assertIn("Translation Review", workbook.sheetnames)
            review_sheet = workbook["Translation Review"]
            self.assertEqual(review_sheet.max_row, 3)
            self.assertEqual(
                [review_sheet.cell(row=row, column=1).value for row in range(2, 4)],
                ["First Lesson (lesson_1.html)", "Second Lesson (lesson_2.html)"],
            )
            self.assertEqual(
                [cell.value for cell in review_sheet[1]],
                ["Page Name", "Type", "Link EN", "Link PT", "Status"],
            )
            self.assertTrue(
                all(
                    review_sheet.cell(row=row, column=3).value is None
                    and review_sheet.cell(row=row, column=4).value is None
                    for row in range(2, 4)
                )
            )
            self.assertEqual(
                [review_sheet.cell(row=row, column=5).value for row in range(2, 4)],
                ["Translated", "Translated"],
            )
        finally:
            shutil.rmtree(hub_dir, ignore_errors=True)

    def test_translation_review_sheet_has_types_dropdown_and_status_colors(self):
        hub_dir = Path.cwd() / "tests" / f"_dashboard_test_{uuid.uuid4().hex}"
        hub_dir.mkdir(parents=True)
        try:
            course_dir = hub_dir / "course_PTBR"
            page_path = course_dir / "wiki_content" / "welcome.html"
            assignment_path = course_dir / "assignment_1" / "assignment_settings.xml"
            quiz_path = course_dir / "quiz_1" / "assessment.xml"
            bank_path = course_dir / "banks" / "question_bank.xml"

            fixtures = {
                page_path: "<html><head><title>Bem-vindo</title></head><body>Texto</body></html>",
                assignment_path: "<assignment><title>Tarefa 1</title></assignment>",
                quiz_path: "<questestinterop><assessment title=\"Teste 1\"/></questestinterop>",
                bank_path: "<questestinterop><objectbank ident=\"bank1\"/></questestinterop>",
            }
            for fixture_path, fixture_content in fixtures.items():
                fixture_path.parent.mkdir(parents=True, exist_ok=True)
                fixture_path.write_text(fixture_content, encoding="utf-8")

            log_path = hub_dir / "translation_log.txt"
            log_lines = [
                "--- New Session (Target: PTBR) ---",
                "[2026-09-16 10:00:00] [System] CourseInfo: Demo Course|DEMO101",
                f"[2026-09-16 10:00:01] [System] TranslatedPage: Welcome | {page_path}",
                f"[2026-09-16 10:00:02] [System] TranslatedPage: Assignment 1 | {assignment_path}",
                f"[2026-09-16 10:00:03] [System] TranslatedPage: Quiz 1 | {quiz_path}",
                f"[2026-09-16 10:00:04] [System] TranslatedPage: Bank 1 | {bank_path}",
            ]
            log_path.write_text("\n".join(log_lines), encoding="utf-8")

            DashboardGenerator(str(log_path), str(hub_dir), "PTBR").generate(lambda _message: None)

            report_path = hub_dir / "Reports" / "Demo Course Translation Report.xlsx"
            workbook = load_workbook(report_path)
            self.assertIn("Translation Review", workbook.sheetnames)

            sheet = workbook["Translation Review"]
            self.assertEqual(
                [cell.value for cell in sheet[1]],
                ["Page Name", "Type", "Link EN", "Link PT", "Status"],
            )
            self.assertEqual(
                [sheet.cell(row=row, column=1).value for row in range(2, 6)],
                [
                    "Welcome (welcome.html)",
                    "Assignment 1 (assignment_settings.xml)",
                    "Quiz 1 (assessment.xml)",
                    "Bank 1 (question_bank.xml)",
                ],
            )
            self.assertEqual(
                [sheet.cell(row=row, column=2).value for row in range(2, 6)],
                ["Page", "Assignment", "Quiz", "Question Bank"],
            )
            self.assertEqual(
                [sheet.cell(row=row, column=5).value for row in range(2, 6)],
                ["Translated"] * 4,
            )
            self.assertTrue(
                all(
                    sheet.cell(row=row, column=3).value is None
                    and sheet.cell(row=row, column=4).value is None
                    for row in range(2, 6)
                )
            )
            self.assertTrue(all(sheet.cell(row=row, column=5).fill.fgColor.rgb.endswith("D9D9D9") for row in range(2, 6)))

            validations = list(sheet.data_validations.dataValidation)
            self.assertEqual(len(validations), 1)
            self.assertEqual(str(validations[0].sqref), "E2:E5")
            self.assertEqual(
                validations[0].formula1,
                '"Translated,Student Reviewed,Professionally Reviewed"',
            )

            conditional_rules = [
                rule
                for conditional_range in sheet.conditional_formatting
                for rule in sheet.conditional_formatting[conditional_range]
            ]
            self.assertEqual(len(conditional_rules), 3)
            self.assertEqual(
                {rule.formula[0] for rule in conditional_rules},
                {'"Translated"', '"Student Reviewed"', '"Professionally Reviewed"'},
            )
            expected_rule_colors = {
                '"Translated"': "FFD9D9D9",
                '"Student Reviewed"': "FFFFF2CC",
                '"Professionally Reviewed"': "FFC6E0B4",
            }
            actual_rule_colors = {
                rule.formula[0]: rule.dxf.fill.bgColor.rgb
                for rule in conditional_rules
            }
            self.assertEqual(actual_rule_colors, expected_rule_colors)
        finally:
            shutil.rmtree(hub_dir, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
