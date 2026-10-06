import shutil
import sys
import unittest
import uuid
from pathlib import Path

from openpyxl import load_workbook


APP_DIR = Path(__file__).resolve().parents[1] / "app"
sys.path.insert(0, str(APP_DIR))

from core.github_report_generator import GitHubTranslationReportGenerator


class GitHubReportGeneratorTests(unittest.TestCase):
    def setUp(self):
        self.root = Path.cwd() / "tests" / f"_github_report_{uuid.uuid4().hex}"
        self.root.mkdir(parents=True)

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_workbook_has_exact_tabs_defaults_and_dropdowns(self):
        path = GitHubTranslationReportGenerator(self.root).generate(
            "course-test1-pt",
            [{"page_name": "Lesson", "type": "Markdown", "status": "Translated"}],
            [
                {"page_title": "One", "context_summary": "input", "variable_en": "name", "variable_pt": "nome", "source_file": "one.md"},
                {"page_title": "Two", "context_summary": "input", "variable_en": "name", "variable_pt": "nome", "source_file": "two.md"},
            ],
            [{
                "asset_path": "images/a.png",
                "package": "External",
                "tl_link": "https://personal-user.github.io/course-test1-pt/",
                "resource_link": "https://personal-user.github.io/course-test1-pt/images/a.png",
                "alt_text_en": "A",
                "alt_text_pt": "Um",
            }],
        )
        workbook = load_workbook(path)
        self.assertEqual(workbook.sheetnames, ["Reviewing", "Variables", "Images and Files"])
        images = workbook["Images and Files"]
        self.assertEqual([images.cell(3, column).value for column in range(1, 12)], [
            "In U.Images", "Images (Copy from UniqueImages)", "Package", "Has Text",
            "Plan", "TL Link", "Resource Link", "Notes", "Num", "Alt Text EN", "Alt Text PT",
        ])
        self.assertIsNone(images["A4"].value)
        self.assertEqual(images["C4"].value, "External")
        self.assertIsNone(images["D4"].value)
        self.assertIsNone(images["E4"].value)
        self.assertEqual(
            images["F4"].hyperlink.target,
            "https://personal-user.github.io/course-test1-pt/",
        )
        self.assertEqual(
            images["G4"].hyperlink.target,
            "https://personal-user.github.io/course-test1-pt/images/a.png",
        )
        formulas = [validation.formula1 for validation in images.data_validations.dataValidation]
        self.assertTrue(any("BrightSpot" in formula and "GitHub" in formula for formula in formulas))
        self.assertTrue(any("Find Equivalent" in formula and "Use Current" in formula for formula in formulas))
        self.assertTrue(workbook["Variables"].conditional_formatting)
