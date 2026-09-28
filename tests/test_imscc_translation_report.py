import shutil
import sys
import threading
import unittest
import uuid
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import load_workbook


APP_DIR = Path(__file__).resolve().parents[1] / "app"
sys.path.insert(0, str(APP_DIR))

from bots.api_utils import call_gemini_with_retry, track_successful_gemini_calls
from bots.xml_bot import XMLTranslationBot
from controller import TranslationController


class SuccessfulGeminiCallTrackingTests(unittest.TestCase):
    def test_notifies_active_translation_context_after_success(self):
        model = Mock()
        model.generate_content.return_value = Mock(text="traduzido")
        callback = Mock()

        with track_successful_gemini_calls(callback):
            response = call_gemini_with_retry(model, "translate", max_retries=1)

        self.assertEqual(response.text, "traduzido")
        callback.assert_called_once_with()

    def test_does_not_notify_when_api_call_fails(self):
        model = Mock()
        model.generate_content.side_effect = RuntimeError("request failed")
        callback = Mock()

        with self.assertRaisesRegex(RuntimeError, "request failed"):
            with track_successful_gemini_calls(callback):
                call_gemini_with_retry(model, "translate", max_retries=1)

        callback.assert_not_called()


class XmlAssessmentTitleTranslationTests(unittest.TestCase):
    def test_translates_assessment_title_attribute(self):
        bot = XMLTranslationBot(api_key="test-key", target_language="PTBR")
        bot._log = Mock()
        bot._translate_batch = Mock(
            side_effect=lambda batch, _constraints: {
                key: (
                    "Questionário do WhatsApp"
                    if value == "W01 Quiz: WhatsApp"
                    else value
                )
                for key, value in batch.items()
            }
        )

        translated = bot.translate_xml_content(
            '<questestinterop><assessment title="W01 Quiz: WhatsApp">'
            '<section/></assessment></questestinterop>'
        )

        self.assertIn('title="Questionário do WhatsApp"', translated)


class ImsccTranslationReportTests(unittest.TestCase):
    def setUp(self):
        self.test_root = Path.cwd() / "tests" / f"_imscc_report_{uuid.uuid4().hex}"
        self.test_root.mkdir(parents=True)

    def tearDown(self):
        shutil.rmtree(self.test_root, ignore_errors=True)

    def test_extracts_quiz_title_from_assessment_attribute(self):
        controller = object.__new__(TranslationController)

        title = controller._extract_page_title(
            '<questestinterop><assessment title="W01 Quiz: WhatsApp ">'
            '</assessment></questestinterop>',
            "xml",
        )

        self.assertEqual(title, "W01 Quiz: WhatsApp")

    def test_review_uses_translated_sibling_title_for_assessment_qti(self):
        quiz_dir = self.test_root / "quiz"
        quiz_dir.mkdir()
        qti_path = quiz_dir / "assessment_qti.xml"
        meta_path = quiz_dir / "assessment_meta.xml"
        qti_path.write_text(
            '<assessment title="W01 Quiz: WhatsApp"><mattext>Sim</mattext></assessment>',
            encoding="utf-8",
        )
        meta_path.write_text(
            '<quiz><title>S01 Questionário: WhatsApp</title>'
            '<description>Descrição</description></quiz>',
            encoding="utf-8",
        )
        controller = object.__new__(TranslationController)
        controller._gemini_translation_pages = {
            str(qti_path): {
                "title": "W01 Quiz: WhatsApp",
                "filepath": str(qti_path),
            }
        }
        controller._gemini_translation_pages_lock = threading.Lock()

        pages = controller.get_gemini_translation_pages()

        self.assertEqual(pages[0]["title"], "S01 Questionário: WhatsApp")

    def test_review_expands_translated_rubric_titles(self):
        rubrics_path = self.test_root / "course_settings" / "rubrics.xml"
        rubrics_path.parent.mkdir()
        rubrics_path.write_text(
            "<rubrics>"
            "<rubric><title>Rubrica da S01</title><description>Critério 1</description></rubric>"
            "<rubric><title>Rubrica da S02</title><description>Critério 2</description></rubric>"
            "</rubrics>",
            encoding="utf-8",
        )
        controller = object.__new__(TranslationController)
        controller._gemini_translation_pages = {
            str(rubrics_path): {
                "title": "Rubrics",
                "filepath": str(rubrics_path),
            }
        }
        controller._gemini_translation_pages_lock = threading.Lock()

        pages = controller.get_gemini_translation_pages()

        self.assertEqual(
            [page["title"] for page in pages],
            ["Rubrica da S01", "Rubrica da S02"],
        )
        self.assertTrue(all(page["filepath"] == str(rubrics_path) for page in pages))

    def test_saved_page_with_successful_gemini_call_is_recorded(self):
        page_path = self.test_root / "wiki_content" / "welcome.html"
        page_path.parent.mkdir(parents=True)
        page_path.write_text(
            "<html><head><title>Welcome</title></head><body>Welcome to the course.</body></html>",
            encoding="utf-8",
        )

        model = Mock()
        model.generate_content.return_value = Mock(
            text=(
                "<html><head><title>Bem-vindo</title></head>"
                "<body>Bem-vindo ao curso.</body></html>"
            )
        )

        class FakeHtmlBot:
            def translate_html_content(self, content, glossary, scriptures, title):
                return call_gemini_with_retry(model, content, max_retries=1).text

        controller = object.__new__(TranslationController)
        controller.target_language = "SPA"
        controller.workspace = Mock(imscc_path="course.imscc")
        controller.workspace.get_target_filepath.side_effect = lambda path: str(path)
        controller.bots = {"html": FakeHtmlBot()}
        controller.link_processor = Mock()
        controller.link_processor.clean_pre_translation_links.side_effect = (
            lambda content, _path, _log: content
        )
        controller.link_processor.rewrite_church_links.side_effect = lambda content: content
        controller.auditor = Mock()
        controller.auditor.get_relevant_terms.return_value = {}
        controller.scripture_checker = Mock()
        controller.scripture_checker.get_scriptures_for_text.return_value = {}
        controller._log = Mock()
        controller._gemini_translation_pages = {}
        controller._gemini_translation_pages_lock = threading.Lock()
        controller._is_already_translated = Mock(return_value=False)

        controller.process_file(str(page_path))

        self.assertEqual(
            controller.get_gemini_translation_pages(),
            [{"title": "Bem-vindo", "filepath": str(page_path.resolve())}],
        )
        self.assertTrue(any(
            "TranslatedPage: Bem-vindo" in call.args[0]
            for call in controller._log.call_args_list
        ))

    def test_imscc_export_is_focused_and_uses_only_tracked_pages(self):
        controller = object.__new__(TranslationController)
        controller.workspace = Mock(imscc_path="course.imscc")
        controller.log_filepath = str(self.test_root / "translation_log.txt")
        controller.hub_dir = str(self.test_root)
        controller.target_language = "PTBR"
        controller._log = Mock()
        controller.get_gemini_translation_pages = Mock(return_value=[
            {"title": "Welcome", "filepath": "course_PTBR/wiki_content/welcome.html"}
        ])

        with patch("controller.DashboardGenerator") as generator_class:
            generator_class.return_value.generate.return_value = "report.xlsx"
            result = controller.update_excel_dashboard()

        self.assertEqual(result, "report.xlsx")
        generator_class.return_value.generate.assert_called_once_with(
            controller._log,
            report_name=None,
            report_code=None,
            review_pages=[
                {"title": "Welcome", "filepath": "course_PTBR/wiki_content/welcome.html"}
            ],
            excluded_sheets=(
                "Dashboard",
                "Bot Analysis",
                "Raw Logs",
                "Translated Pages",
                "Link Actions",
                "Missing Alt Texts",
            ),
        )

    def test_review_candidates_exclude_manifest_qti_and_title_only_xml(self):
        html_path = self.test_root / "page.html"
        setup_notes_path = self.test_root / "setup-notes-and-course-settings-2.html"
        text_path = self.test_root / "notes.txt"
        qti_path = self.test_root / "quiz.qti"
        bank_qti_path = self.test_root / "question_bank.qti"
        canvas_export_path = self.test_root / "canvas_export.txt"
        manifest_path = self.test_root / "imsmanifest.xml"
        lti_path = self.test_root / "tutoring.xml"
        title_only_path = self.test_root / "assignment_settings.xml"
        content_xml_path = self.test_root / "assessment_meta.xml"

        html_path.write_text("<p>Page content</p>", encoding="utf-8")
        setup_notes_path.write_text("<p>Setup instructions</p>", encoding="utf-8")
        text_path.write_text("Text content", encoding="utf-8")
        qti_path.write_text("<assessment><mattext>Question</mattext></assessment>", encoding="utf-8")
        bank_qti_path.write_text(
            "<questestinterop><objectbank>"
            "<qtimetadatafield><fieldlabel>bank_title</fieldlabel>"
            "<fieldentry>Banco de perguntas</fieldentry></qtimetadatafield>"
            "<item title=\"Pergunta 1\"><mattext>Texto da pergunta</mattext></item>"
            "</objectbank></questestinterop>",
            encoding="utf-8",
        )
        canvas_export_path.write_text("Canvas export marker", encoding="utf-8")
        manifest_path.write_text("<manifest><description>Copyright</description></manifest>", encoding="utf-8")
        lti_path.write_text(
            "<cartridge_basiclti_link><title>Tutoring</title>"
            "<description>External navigation link.</description>"
            "</cartridge_basiclti_link>",
            encoding="utf-8",
        )
        title_only_path.write_text("<assignment><title>Assignment name</title></assignment>", encoding="utf-8")
        content_xml_path.write_text(
            "<quiz><title>Quiz name</title><description>Study the lesson.</description></quiz>",
            encoding="utf-8",
        )

        self.assertTrue(TranslationController._is_translation_review_candidate(str(html_path)))
        self.assertFalse(TranslationController._is_translation_review_candidate(str(setup_notes_path)))
        self.assertTrue(TranslationController._is_translation_review_candidate(str(text_path)))
        self.assertTrue(TranslationController._is_translation_review_candidate(str(content_xml_path)))
        self.assertFalse(TranslationController._is_translation_review_candidate(str(qti_path)))
        self.assertTrue(TranslationController._is_translation_review_candidate(str(bank_qti_path)))
        self.assertFalse(TranslationController._is_translation_review_candidate(str(canvas_export_path)))
        self.assertFalse(TranslationController._is_translation_review_candidate(str(manifest_path)))
        self.assertFalse(TranslationController._is_translation_review_candidate(str(lti_path)))
        self.assertFalse(TranslationController._is_translation_review_candidate(str(title_only_path)))

    def test_imscc_workbook_matches_single_sheet_edtech_export(self):
        page_path = self.test_root / "course_PTBR" / "wiki_content" / "welcome.html"
        page_path.parent.mkdir(parents=True)
        page_path.write_text("<title>Bem-vindo</title>", encoding="utf-8")
        log_path = self.test_root / "translation_log.txt"
        log_path.write_text(
            "\n".join([
                "--- New Session (Target: PTBR) ---",
                "[2026-09-28 10:00:00] [System] CourseInfo: IMSCC Demo|DEMO101",
            ]),
            encoding="utf-8",
        )

        controller = object.__new__(TranslationController)
        controller.workspace = Mock(imscc_path="course.imscc")
        controller.log_filepath = str(log_path)
        controller.hub_dir = str(self.test_root)
        controller.target_language = "PTBR"
        controller._log = Mock()
        controller.get_gemini_translation_pages = Mock(return_value=[
            {"title": "Welcome", "filepath": str(page_path)}
        ])

        report_path = controller.update_excel_dashboard()

        workbook = load_workbook(report_path)
        self.assertEqual(workbook.sheetnames, ["Translation Review"])
        sheet = workbook["Translation Review"]
        self.assertEqual(sheet["A2"].value, "Welcome (welcome.html)")
        self.assertEqual(sheet["B2"].value, "Page")
        self.assertEqual(sheet["E2"].value, "Translated")


if __name__ == "__main__":
    unittest.main()
